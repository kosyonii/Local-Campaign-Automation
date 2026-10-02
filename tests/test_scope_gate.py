"""Gate A(비모바일 scope) 단위 테스트. Gemini는 가짜 judge로 대체한다."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from quality_control import run as qc_run  # noqa: E402
from quality_control.schema import (  # noqa: E402
    RAW_SHEET_ORIGINAL,
    RAW_SHEET_SUBSIDIARY,
    GateContext,
    GateOutcome,
    RawRow,
    Verdict,
    merge_outcomes,
)
from quality_control.scope_gate import (  # noqa: E402
    DEFAULT_SCOPE_CONFIG_PATH,
    Judgement,
    ScopeQuery,
    load_scope_dictionary,
    run_scope_gate,
    scan_text,
)

SUBSIDIARY_COLUMNS = [
    "Conversation Stream",
    "Campaign ID",
    "Profile URL",
    "User Name",
    "Permalink",
    "Created Time",
    "snType column",
    "Author Screen Name",
    "Media Type",
    "Media URL",
    "Sender Screen Name",
    "Sender Follower Count",
    "Sender Bio",
]


@pytest.fixture(scope="module")
def dictionary():
    return load_scope_dictionary(DEFAULT_SCOPE_CONFIG_PATH)


def row(text: str, number: int = 2, sheet: str = RAW_SHEET_SUBSIDIARY, **extra):
    values = {"Conversation Stream": text, **extra}
    return RawRow(sheet=sheet, row_number=number, values=values)


def ctx(tmp_path: Path, use_llm: bool = True) -> GateContext:
    source = tmp_path / "in.xlsx"
    source.write_bytes(b"x")
    return GateContext(
        input_path=source,
        output_path=tmp_path / "out.xlsx",
        use_llm=use_llm,
        workers=2,
    )


def non_mobile(evidence: str, products=("TV",), mobile=(), confidence=95):
    return Judgement(
        verdict="NON_MOBILE_ONLY",
        evidence=evidence,
        non_mobile_products=tuple(products),
        mobile_products=tuple(mobile),
        reason="TV 홍보",
        confidence=confidence,
    )


def unclear():
    return Judgement("UNCLEAR", "", (), (), "매체명으로 쓰인 TV", 40)


# --- 사전 / 매칭 ------------------------------------------------------------


@pytest.mark.parametrize(
    "text,family",
    [
        ("New #TV drop", "TV"),
        ("SamsungTV deals", "TV"),
        ("Samsung_TV fans", "TV"),
        ("TV를 보세요", "TV"),
        ("ＴＶ全角", "TV"),
        ("Odyssey G9 gaming", "Odyssey"),
        ("#SamsungSoundbar", "Soundbar"),
        ("#SamsungBespoke", "Bespoke"),
        ("Smartswitch tips", "SmartSwitch"),
        ("#Smart_Switch", "SmartSwitch"),
        ("Samsungcare+ plan", "SamsungCare+"),
        ("삼성케어플러스 가입", "SamsungCare+"),
        ("Soundbar sale", "Soundbar"),
        ("JBL party", "JBL"),
        ("비스포크 냉장고", "Bespoke"),
    ],
)
def test_nonmobile_terms_hit(dictionary, text, family):
    assert family in scan_text(text, dictionary).nonmobile


@pytest.mark.parametrize(
    "text",
    [
        "Molotov cocktail",  # tv가 단어 안에 없음
        "iptv channel",
        "TV5 Monde",  # 숫자 인접
        "Galaxy S25 launch",
        "Just a normal post",
        "",
    ],
)
def test_short_token_boundary_no_false_hit(dictionary, text):
    assert not scan_text(text, dictionary).nonmobile


@pytest.mark.parametrize(
    "text",
    [
        "Galaxy Z Fold8 with Smart Switch",
        "갤럭시 Z플립7",
        "Galaxy Watch8 #SmartSwitch",
        "Galaxy Buds3 #JBL",
        "S25 Ultra camera",
        "Z Flip7 now",
        "Fold8 is here",
    ],
)
def test_strong_mobile_signals(dictionary, text):
    assert scan_text(text, dictionary).strong


def test_weak_signal_is_not_strong(dictionary):
    scan = scan_text("Dein Phone, deine Size #Bespoke", dictionary)
    assert scan.weak and not scan.strong


# --- 게이트 판정 --------------------------------------------------------------


def run_gate(rows, tmp_path, judge, use_llm=True):
    return run_scope_gate(rows, ctx(tmp_path, use_llm), judge=judge)


def never_called(query: ScopeQuery):
    raise AssertionError(f"LLM이 호출되면 안 됨: {query.key}")


def test_no_keyword_is_untouched(tmp_path):
    out = run_gate([row("Galaxy S25 giveaway")], tmp_path, never_called)
    assert out == {}


def test_strong_mobile_keeps_without_llm(tmp_path):
    rows = [row("Smart Switch로 새 폰 설정 #GalaxyZFold8 갤럭시 Z Fold8")]
    out = run_gate(rows, tmp_path, never_called)
    outcome = out[rows[0].key]
    assert outcome.verdict is Verdict.KEEP
    assert outcome.rule == "SCOPE_NONMOBILE_KW_WITH_MOBILE"


def test_weak_mobile_flags_without_llm(tmp_path):
    rows = [row("Dein Phone, deine Size #Bespoke")]
    out = run_gate(rows, tmp_path, never_called)
    assert out[rows[0].key].verdict is Verdict.FLAG
    assert out[rows[0].key].rule == "SCOPE_NONMOBILE_KW_WEAK_MOBILE"


def test_pure_nonmobile_with_verified_evidence_drops(tmp_path):
    text = "Discover the new Samsung TV lineup with Neo QLED"
    rows = [row(text)]
    out = run_gate(
        rows,
        tmp_path,
        lambda q: non_mobile("new Samsung TV lineup"),
    )
    outcome = out[rows[0].key]
    assert outcome.verdict is Verdict.DROP
    assert outcome.rule == "SCOPE_NONMOBILE_ONLY"
    assert outcome.confidence == 95  # 기록만 되고 판정에는 쓰이지 않음


def test_evidence_check_ignores_case_and_whitespace(tmp_path):
    rows = [row("Great   Odyssey\nG9 deal")]
    out = run_gate(
        rows,
        tmp_path,
        lambda q: non_mobile("great odyssey g9", ("Odyssey G9",)),
    )
    assert out[rows[0].key].verdict is Verdict.DROP


def test_hallucinated_evidence_is_flagged_not_dropped(tmp_path):
    rows = [row("Discover the new Samsung TV lineup")]
    out = run_gate(
        rows,
        tmp_path,
        lambda q: non_mobile("a sentence that is not in the post"),
    )
    assert out[rows[0].key].verdict is Verdict.FLAG
    assert out[rows[0].key].rule == "SCOPE_DROP_UNVERIFIED"


def test_high_confidence_never_overrides_failed_checks(tmp_path):
    rows = [row("Discover the new Samsung TV lineup")]
    out = run_gate(
        rows,
        tmp_path,
        lambda q: non_mobile("new Samsung TV", mobile=("Galaxy S25",), confidence=100),
    )
    assert out[rows[0].key].verdict is Verdict.FLAG


def test_missing_product_blocks_drop(tmp_path):
    rows = [row("Discover the new Samsung TV lineup")]
    out = run_gate(
        rows,
        tmp_path,
        lambda q: non_mobile("new Samsung TV", products=()),
    )
    assert out[rows[0].key].verdict is Verdict.FLAG


def test_tv_as_media_name_is_flagged(tmp_path):
    rows = [row("Molotov TV Awards 2026 ✨ merci à tous")]
    out = run_gate(rows, tmp_path, lambda q: unclear())
    assert out[rows[0].key].verdict is Verdict.FLAG
    assert out[rows[0].key].rule == "SCOPE_LLM_UNCLEAR"


def test_llm_mobile_verdict_is_flagged(tmp_path):
    rows = [row("TV time with friends")]
    out = run_gate(
        rows,
        tmp_path,
        lambda q: Judgement("MOBILE_RELATED", "", (), ("Galaxy",), "폰 맥락", 70),
    )
    assert out[rows[0].key].verdict is Verdict.FLAG
    assert out[rows[0].key].rule == "SCOPE_LLM_MOBILE"


def test_llm_error_is_flagged(tmp_path):
    def broken(query):
        raise RuntimeError("503 unavailable")

    rows = [row("New Samsung TV")]
    out = run_gate(rows, tmp_path, broken)
    assert out[rows[0].key].verdict is Verdict.FLAG
    assert out[rows[0].key].rule == "SCOPE_LLM_UNAVAILABLE"
    assert "503" in out[rows[0].key].reason


def test_no_llm_mode_never_drops(tmp_path):
    rows = [row("New Samsung TV")]
    out = run_gate(rows, tmp_path, never_called, use_llm=False)
    assert out[rows[0].key].verdict is Verdict.FLAG
    assert out[rows[0].key].rule == "SCOPE_LLM_UNAVAILABLE"


def test_original_sheet_rows_are_not_evaluated(tmp_path):
    rows = [row("New Samsung TV", sheet=RAW_SHEET_ORIGINAL)]
    assert run_gate(rows, tmp_path, never_called) == {}


def test_account_name_counts_but_url_does_not(tmp_path):
    by_name = row("Big sale today", **{"Author Screen Name": "SamsungTV_UK"})
    by_url = row(
        "Big sale today",
        number=3,
        **{"Permalink": "https://instagram.com/tv/abc", "Profile URL": "https://x.com/tv_fan"},
    )
    out = run_gate(
        [by_name, by_url],
        tmp_path,
        lambda q: non_mobile("Big sale today", ("Samsung TV",)),
    )
    assert by_name.key in out
    assert by_url.key not in out


# --- 병합 --------------------------------------------------------------------


def test_merge_takes_heaviest_verdict_and_all_rules():
    merged = merge_outcomes(
        [
            GateOutcome(Verdict.FLAG, "A", "a", 50),
            GateOutcome(Verdict.DROP, "B", "b", 90, "C1"),
            GateOutcome(Verdict.KEEP, "C", "c"),
        ]
    )
    assert merged.verdict is Verdict.DROP
    assert merged.rule == "A | B | C"
    assert merged.confidence == 90
    assert merged.cluster_id == "C1"
    assert "[B] b" in merged.reason


# --- 엔드투엔드 ----------------------------------------------------------------


def build_input(path: Path):
    wb = Workbook()
    original = wb.active
    original.title = RAW_SHEET_ORIGINAL
    original.append(["Conversation Stream", "Campaign ID", "Sender Follower Count"])
    original.append(["Partner TV post", "p1", 1000])

    sub = wb.create_sheet(RAW_SHEET_SUBSIDIARY)
    sub.append(SUBSIDIARY_COLUMNS)

    def add(text):
        sub.append([text, "c1", None, "acct", None, None, None, "acct"] + [None] * 5)

    add("Galaxy S25 giveaway")  # row 2 KEEP
    add("Discover the new Samsung TV lineup")  # row 3 DROP
    add("Dein Phone, deine Size #Bespoke")  # row 4 FLAG
    add("Molotov TV Awards merci")  # row 5 FLAG
    add("Smart Switch 갤럭시 Z Fold8")  # row 6 KEEP(규칙 기록)
    wb.save(path)


def fake_judge(query: ScopeQuery) -> Judgement:
    if "Samsung TV" in query.text:
        return non_mobile("new Samsung TV lineup")
    return unclear()


def test_end_to_end_workbook(tmp_path, monkeypatch):
    source = tmp_path / "raw.xlsx"
    build_input(source)
    before = source.read_bytes()

    monkeypatch.setattr("quality_control.scope_gate.build_gemini_judge", lambda: fake_judge)

    out = tmp_path / "raw_qc.xlsx"
    assert qc_run.main([str(source), "-o", str(out), "--workers", "2"]) == 0

    assert source.read_bytes() == before  # 원본 미변경

    wb = load_workbook(out)
    assert wb.sheetnames == [
        RAW_SHEET_ORIGINAL,
        RAW_SHEET_SUBSIDIARY,
        "QC_Full",
        "QC_Clean",
        "QC_Dropped",
    ]
    # 원본 시트 보존
    assert wb[RAW_SHEET_ORIGINAL].max_row == 2
    assert wb[RAW_SHEET_SUBSIDIARY].max_row == 6
    assert wb[RAW_SHEET_SUBSIDIARY].cell(1, 1).value == "Conversation Stream"

    def table(name):
        sheet = wb[name]
        header = [c.value for c in sheet[1]]
        return [dict(zip(header, r)) for r in sheet.iter_rows(min_row=2, values_only=True)]

    flagged, clean, dropped = table("QC_Full"), table("QC_Clean"), table("QC_Dropped")

    assert len(flagged) == 6  # 원문 1 + 전략법인 5
    assert len(dropped) == 1 and len(clean) == 5
    assert dropped[0]["QC_Source_Row"] == 3
    assert dropped[0]["QC_Verdict"] == "DROP"
    assert dropped[0]["QC_Rule"] == "SCOPE_NONMOBILE_ONLY"
    assert dropped[0]["QC_Confidence"] == 95
    assert "Samsung TV" in dropped[0]["QC_Reason"]
    assert all(r["QC_Verdict"] != "DROP" for r in clean)

    verdicts = {(r["QC_Source_Sheet"], r["QC_Source_Row"]): r["QC_Verdict"] for r in flagged}
    assert verdicts[(RAW_SHEET_ORIGINAL, 2)] == "KEEP"
    assert verdicts[(RAW_SHEET_SUBSIDIARY, 2)] == "KEEP"
    assert verdicts[(RAW_SHEET_SUBSIDIARY, 4)] == "FLAG"
    assert verdicts[(RAW_SHEET_SUBSIDIARY, 5)] == "FLAG"
    assert verdicts[(RAW_SHEET_SUBSIDIARY, 6)] == "KEEP"

    # 원문 시트에만 있는 컬럼과 전략법인 전용 컬럼이 합집합으로 모두 존재
    assert "Author Screen Name" in flagged[0]
    assert flagged[0]["Campaign ID"] == "p1"

    # 체크포인트 정리됨
    assert not list(tmp_path.glob(".*.checkpoint.json"))


def test_rerun_on_qc_output_is_idempotent(tmp_path, monkeypatch):
    source = tmp_path / "raw.xlsx"
    build_input(source)
    monkeypatch.setattr("quality_control.scope_gate.build_gemini_judge", lambda: fake_judge)

    first = tmp_path / "raw_qc.xlsx"
    qc_run.main([str(source), "-o", str(first), "--workers", "2"])
    second = tmp_path / "again.xlsx"
    qc_run.main([str(first), "-o", str(second), "--workers", "2"])

    wb = load_workbook(second)
    assert wb.sheetnames.count("QC_Full") == 1
    assert wb["QC_Full"].max_row == 7
    assert wb["QC_Dropped"].max_row == 2


def test_refuses_to_overwrite_input(tmp_path):
    source = tmp_path / "raw.xlsx"
    build_input(source)
    assert qc_run.main([str(source), "-o", str(source), "--no-llm"]) == 2


def test_default_output_name(tmp_path):
    assert qc_run.default_output_path(tmp_path / "a.xlsx").name == "a_qc.xlsx"
    assert qc_run.default_output_path(tmp_path / "a_qc.xlsx").name == "a_qc.xlsx"


def test_no_llm_cli_flags_instead_of_dropping(tmp_path):
    source = tmp_path / "raw.xlsx"
    build_input(source)
    out = tmp_path / "raw_qc.xlsx"
    assert qc_run.main([str(source), "-o", str(out), "--no-llm"]) == 0
    wb = load_workbook(out)
    assert wb["QC_Dropped"].max_row == 1  # 헤더만
    assert wb["QC_Clean"].max_row == 7


# --- 체크포인트 ----------------------------------------------------------------


def test_checkpoint_resume_skips_successful_rows(tmp_path):
    source = tmp_path / "raw.xlsx"
    build_input(source)
    out = tmp_path / "raw_qc.xlsx"
    calls: list[str] = []
    fail = {"Molotov"}

    def flaky(query: ScopeQuery) -> Judgement:
        calls.append(query.key)
        if any(word in query.text for word in fail):
            raise RuntimeError("503 unavailable")
        return fake_judge(query)

    from quality_control import scope_gate

    wb = load_workbook(source)
    rows, _ = qc_run.read_raw_rows(wb)
    context = GateContext(source, out, use_llm=True, workers=1)

    first = scope_gate.run_scope_gate(rows, context, judge=flaky)
    assert first[f"{RAW_SHEET_SUBSIDIARY}!3"].verdict is Verdict.DROP
    assert first[f"{RAW_SHEET_SUBSIDIARY}!5"].rule == "SCOPE_LLM_UNAVAILABLE"
    assert len(calls) == 2
    assert context.extra["scope_checkpoint_path"].is_file()  # 오류 시 저장됨

    calls.clear()
    fail.clear()
    context2 = GateContext(source, out, use_llm=True, workers=1)
    second = scope_gate.run_scope_gate(rows, context2, judge=flaky)

    assert calls == [f"{RAW_SHEET_SUBSIDIARY}!5"]  # 성공했던 3행은 재호출 안 함
    assert second[f"{RAW_SHEET_SUBSIDIARY}!3"].verdict is Verdict.DROP
    assert second[f"{RAW_SHEET_SUBSIDIARY}!5"].rule == "SCOPE_LLM_UNCLEAR"


def test_checkpoint_ignored_when_input_changes(tmp_path):
    source = tmp_path / "raw.xlsx"
    build_input(source)
    out = tmp_path / "raw_qc.xlsx"
    from quality_control import scope_gate

    wb = load_workbook(source)
    rows, _ = qc_run.read_raw_rows(wb)
    context = GateContext(source, out, use_llm=True, workers=1)

    def boom(query):
        raise RuntimeError("503")

    # 일부 성공 + 일부 실패 상태의 체크포인트를 만든다
    scope_gate.run_scope_gate(
        rows, context, judge=lambda q: fake_judge(q) if "Samsung" in q.text else boom(q)
    )

    build_input(source)  # 같은 내용이어도 mtime/size 변경 -> fingerprint 변경
    import os, time

    os.utime(source, (time.time() + 5, time.time() + 5))

    calls: list[str] = []

    def counting(query):
        calls.append(query.key)
        return fake_judge(query)

    scope_gate.run_scope_gate(
        rows, GateContext(source, out, use_llm=True, workers=1), judge=counting
    )
    assert len(calls) == 2  # 전부 다시 호출
