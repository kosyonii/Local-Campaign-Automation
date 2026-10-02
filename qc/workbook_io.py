"""Raw Excel 읽기와 QC 시트(QC_Flagged / QC_Clean / QC_Dropped) 쓰기.

- 원본 파일은 건드리지 않고 별도 파일로 저장한다. 원본 시트는 그대로 보존된다.
- 이미 QC 시트가 있는 파일을 다시 처리하면 QC 시트만 새로 만든다(재실행 멱등).
- 원본에 이미지/차트가 있는 formatted 파일은 openpyxl이 보존하지 못할 수 있다.
  이 도구의 입력은 1단계 직후의 raw 파일이다.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .schema import (
    KEEP_RECORD,
    QC_COLUMNS,
    QC_SHEET_CLEAN,
    QC_SHEET_DROPPED,
    QC_SHEET_FLAGGED,
    QC_SHEETS,
    QC_TRACE_COLUMNS,
    RAW_SHEETS,
    QCRecord,
    RawRow,
    Verdict,
)

HEADER_FILL = PatternFill("solid", start_color="D9D9D9")
VERDICT_FILLS = {
    Verdict.DROP.value: PatternFill("solid", start_color="F4B6B6"),
    Verdict.FLAG.value: PatternFill("solid", start_color="FFE699"),
}


def load_input_workbook(path: Path) -> Workbook:
    return load_workbook(path)


def read_raw_rows(
    workbook: Workbook,
) -> tuple[list[RawRow], list[str]]:
    """Raw 시트들의 데이터 행과, 시트 컬럼의 합집합(순서 유지)을 반환한다."""

    rows: list[RawRow] = []
    columns: list[str] = []
    found_any = False

    # 컬럼이 더 많은 전략법인 시트를 기준으로 합집합을 만든다.
    for sheet_name in reversed(RAW_SHEETS):
        if sheet_name not in workbook.sheetnames:
            continue

        found_any = True
        sheet = workbook[sheet_name]
        header = [
            None if cell is None else str(cell).strip()
            for cell in next(
                sheet.iter_rows(min_row=1, max_row=1, values_only=True),
                (),
            )
        ]

        if not any(header):
            raise ValueError(f"{sheet_name}: 1행에 헤더가 없습니다.")

        for name in header:
            if name and name not in columns:
                columns.append(name)

        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if all(v is None or v == "" for v in values):
                continue

            rows.append(
                RawRow(
                    sheet=sheet_name,
                    row_number=row_number,
                    values={
                        name: value
                        for name, value in zip(header, values)
                        if name
                    },
                )
            )

    if not found_any:
        raise ValueError(
            f"Raw 시트를 찾지 못했습니다: {', '.join(RAW_SHEETS)}"
        )

    # 시트 순서(원문 → 전략법인), 시트 안에서는 행 순서로 정렬
    order = {name: i for i, name in enumerate(RAW_SHEETS)}
    rows.sort(key=lambda r: (order[r.sheet], r.row_number))

    return rows, columns


def _write_sheet(
    workbook: Workbook,
    title: str,
    columns: list[str],
    entries: list[tuple[RawRow, QCRecord]],
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(columns)

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL

    verdict_col = columns.index("QC_Verdict") + 1

    for raw, record in entries:
        cells: dict[str, Any] = dict(raw.values)
        cells["QC_Source_Sheet"] = raw.sheet
        cells["QC_Source_Row"] = raw.row_number
        cells.update(record.as_cells())

        sheet.append([cells.get(name) for name in columns])

        fill = VERDICT_FILLS.get(record.verdict.value)
        if fill is not None:
            sheet.cell(
                row=sheet.max_row, column=verdict_col
            ).fill = fill

    sheet.freeze_panes = "A2"

    if entries:
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(columns))}{sheet.max_row}"
        )

    for index, name in enumerate(columns, start=1):
        width = 60 if name in ("QC_Reason", "Conversation Stream") else 18
        sheet.column_dimensions[get_column_letter(index)].width = width

    reason_col = columns.index("QC_Reason") + 1
    for row in sheet.iter_rows(min_row=2, min_col=reason_col, max_col=reason_col):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_qc_workbook(
    workbook: Workbook,
    output_path: Path,
    rows: list[RawRow],
    raw_columns: list[str],
    records: dict[str, QCRecord],
) -> dict[str, int]:
    """QC 시트 3개를 만들고 output_path에 저장한다. 시트별 행 수를 반환."""

    for name in QC_SHEETS:
        if name in workbook.sheetnames:
            del workbook[name]

    columns = [*raw_columns, *QC_TRACE_COLUMNS, *QC_COLUMNS]
    entries = [(r, records.get(r.key, KEEP_RECORD)) for r in rows]

    clean = [e for e in entries if e[1].verdict is not Verdict.DROP]
    dropped = [e for e in entries if e[1].verdict is Verdict.DROP]

    _write_sheet(workbook, QC_SHEET_FLAGGED, columns, entries)
    _write_sheet(workbook, QC_SHEET_CLEAN, columns, clean)
    _write_sheet(workbook, QC_SHEET_DROPPED, columns, dropped)

    workbook.save(output_path)

    return {
        QC_SHEET_FLAGGED: len(entries),
        QC_SHEET_CLEAN: len(clean),
        QC_SHEET_DROPPED: len(dropped),
    }
