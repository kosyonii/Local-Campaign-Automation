# ===============================================================================================
# Formatted Excel -> Consumer Reaction URL 추출
#
# 주요 처리:
# 1. raw_to_processed.py(2a)가 만든 *_formatted.xlsx의
#    "로컬 캠페인 리스트_QHB8" 시트를 읽는다 (Permalink만 있는 URL 컬럼).
# 2. 채널별로 소비자 댓글/reply URL을 추출해 같은 URL 컬럼을
#    "POST_URL\nCOMMENT_URL" 형식으로 덮어쓴다.
# 3. Excel의 다른 컬럼(정제 결과)은 건드리지 않는다.
#
# 경로:
# raw_to_processed.py와 동일하게 기존 실행 output 폴더를 사용한다.
# 이 모듈은 새로운 차수 폴더를 생성하지 않는다.
# ===============================================================================================

# =========================================================
# 0. Import library & Set variables
# =========================================================

import argparse
import json
import os
import time
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from comment_extractor_v2 import (
    CommentExtractorSession,
    build_url_cell_value,
    INSTAGRAM_PREFETCH_MAX_WORKERS,
    YOUTUBE_PREFETCH_MAX_WORKERS,
)
from raw_to_processed import (
    NEW_SHEET_NAME,
    OUTPUT_COLUMNS,
    ENV_INPUT_DATE,
    ENV_RUN_NUMBER,
    ENV_OUTPUT_DIR,
    find_header_row,
    find_column_index,
    get_output_column_index,
    normalize_url,
    resolve_output_directory,
    resolve_path_from_project,
    wait_for_network_ready,
)


BASE_DIR = Path(__file__).resolve().parent

# 메모리 부족 등으로 프로세스가 외부에서 강제 종료되는 경우를 대비해
# 워크북 + 진행 상태를 주기적으로 저장하고, 재실행 시 그 지점부터
# 이어서 처리한다 (처음부터 다시 돌리지 않기 위함).
CHECKPOINT_INTERVAL_SECONDS = 180.0


# =========================================================
# 1. Excel 경로 결정
# =========================================================

def build_formatted_excel_path(
    input_date: str,
    output_dir: Path,
) -> Path:
    """
    2a(raw_to_processed.py)가 만든 formatted Excel 경로를 반환한다.

    입력이자 출력:
        {실행 output 폴더}/
        {YYMMDD}_SLCC_SOV_Local Campaign Tracking_{월}월_v01_formatted.xlsx
    """

    try:
        input_date_obj = datetime.strptime(
            input_date,
            "%y%m%d",
        )
    except ValueError as exc:
        raise ValueError(
            "날짜는 YYMMDD 형식으로 입력해야 합니다. "
            "예시: 260724"
        ) from exc

    input_month = input_date_obj.month

    formatted_excel_path = output_dir / (
        f"{input_date}_SLCC_SOV_Local Campaign Tracking_"
        f"{input_month}월_v01_formatted.xlsx"
    )

    if not formatted_excel_path.is_file():
        raise FileNotFoundError(
            "Formatted Excel file not found (2a 단계가 먼저 완료되어야 합니다): "
            f"{formatted_excel_path}"
        )

    return formatted_excel_path


def prepare_temporary_path(
    target_path: Path,
) -> Path:
    """
    같은 파일을 직접 덮어쓰지 않고 임시 파일 경로를 준비한다.
    """

    temporary_path = target_path.with_name(
        f".{target_path.stem}.reaction.partial.xlsx"
    )

    if temporary_path.exists():
        temporary_path.unlink()

    return temporary_path


def cleanup_temporary_path(
    temporary_path: Path,
) -> None:
    if temporary_path.exists():
        temporary_path.unlink()


def source_fingerprint(
    path: Path,
) -> str:
    """
    체크포인트가 어떤 formatted Excel을 기준으로 만들어졌는지 식별한다.
    2a가 formatted Excel을 새로 만들면 (수정 시각/크기가 바뀌면)
    이전 체크포인트는 옛 데이터 기준이므로 재사용하면 안 된다.
    """

    stat = path.stat()

    return f"{stat.st_mtime_ns}:{stat.st_size}"


def extract_post_url(
    raw_url,
) -> str | None:
    """
    URL 컬럼 값에서 게시물 URL만 꺼낸다.

    2b가 이미 처리한 파일을 다시 처리하는 경우 셀 값이
    "POST_URL(줄바꿈)COMMENT_URL" 형식이므로 첫 줄만 게시물 URL로 사용한다.
    """

    if raw_url is None:
        return None

    first_line = str(raw_url).strip().split("\n", 1)[0]

    return normalize_url(first_line)


def checkpoint_paths(
    formatted_excel_path: Path,
) -> tuple[Path, Path]:
    """
    중간 저장(체크포인트) 워크북 및 진행 상태 JSON 경로를 반환한다.
    """

    checkpoint_xlsx_path = formatted_excel_path.with_name(
        f".{formatted_excel_path.stem}.reaction_checkpoint.xlsx"
    )
    checkpoint_json_path = formatted_excel_path.with_name(
        f".{formatted_excel_path.stem}.reaction_checkpoint.json"
    )

    return checkpoint_xlsx_path, checkpoint_json_path


def load_checkpoint(
    checkpoint_xlsx_path: Path,
    checkpoint_json_path: Path,
    expected_fingerprint: str,
):
    """
    이전 실행이 남긴 체크포인트가 있으면 불러온다.

    둘 중 하나라도 없거나 손상되어 읽을 수 없거나, 체크포인트가 다른
    (이전) formatted Excel 기준으로 만들어졌으면 (None, None)을 반환해
    처음부터 새로 시작하게 한다.
    """

    if (
        not checkpoint_xlsx_path.is_file()
        or not checkpoint_json_path.is_file()
    ):
        return None, None

    try:
        progress = json.loads(
            checkpoint_json_path.read_text(encoding="utf-8")
        )
        workbook = load_workbook(checkpoint_xlsx_path)
    except Exception as exc:
        print(
            "[WARNING] 체크포인트를 읽지 못해 "
            "처음부터 다시 시작합니다: "
            f"{type(exc).__name__}: {exc}"
        )
        return None, None

    if progress.get("source_fingerprint") != expected_fingerprint:
        print(
            "[WARNING] 체크포인트가 현재 formatted Excel과 "
            "다른 파일 기준으로 만들어져 있어(2a가 새로 생성했거나 "
            "옛 형식) 무시하고 처음부터 시작합니다."
        )
        return None, None

    return workbook, progress


def make_checkpoint_saver(
    workbook,
    checkpoint_xlsx_path: Path,
    checkpoint_json_path: Path,
    fingerprint: str,
    interval_seconds: float = CHECKPOINT_INTERVAL_SECONDS,
):
    """
    workbook + 진행 상태를 주기적으로(interval_seconds마다) 원자적으로
    저장하는 콜백을 만든다. 임시 파일에 저장 후 os.replace로 교체하므로
    저장 도중 프로세스가 죽어도 기존 체크포인트는 손상되지 않는다.
    """

    state = {"last_saved_at": 0.0}

    def save(
        *,
        next_row: int | None,
        processed_count: int,
        reaction_found_count: int,
        force: bool = False,
    ) -> None:
        now = time.monotonic()

        if (
            not force
            and (now - state["last_saved_at"]) < interval_seconds
        ):
            return

        temp_xlsx_path = checkpoint_xlsx_path.with_suffix(".tmp.xlsx")
        workbook.save(temp_xlsx_path)
        os.replace(temp_xlsx_path, checkpoint_xlsx_path)

        payload = {
            "source_fingerprint": fingerprint,
            "next_row": next_row,
            "processed_count": processed_count,
            "reaction_found_count": reaction_found_count,
            "saved_at": datetime.now(timezone.utc).isoformat(),
        }

        temp_json_path = checkpoint_json_path.with_suffix(".tmp.json")
        temp_json_path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        os.replace(temp_json_path, checkpoint_json_path)

        state["last_saved_at"] = now

        print(
            "[CHECKPOINT] 저장 완료 "
            f"(next_row={next_row}, processed={processed_count})"
        )

    return save


def cleanup_checkpoint(
    checkpoint_xlsx_path: Path,
    checkpoint_json_path: Path,
) -> None:
    """
    정상 완료 후 더 이상 필요 없는 체크포인트 파일을 정리한다.
    """

    for path in (checkpoint_xlsx_path, checkpoint_json_path):
        if path.exists():
            path.unlink()


# =========================================================
# 2. 시트 처리
# =========================================================

def process_reaction_sheet(
    ws,
    comment_session: CommentExtractorSession,
    start_row: int | None = None,
    start_processed_count: int = 0,
    start_reaction_found_count: int = 0,
    checkpoint_saver=None,
) -> tuple[int, int]:
    """
    로컬 캠페인 리스트_QHB8 시트의 각 행을 순회하며
    URL 컬럼을 POST_URL\\nCOMMENT_URL 형식으로 덮어쓴다.

    start_row가 주어지면 (체크포인트에서 재개하는 경우)
    해당 행부터 처리를 시작하고, 이전까지의 카운트를 이어받는다.

    반환값: (처리한 행 수, 소비자 반응 URL을 찾은 행 수)
    """

    required_columns = {
        "Channel",
        "URL",
    }

    header_row = find_header_row(
        ws=ws,
        required_columns=required_columns,
    )

    if header_row is None:
        raise ValueError(
            f"{ws.title}: required header row not found "
            f"(columns={required_columns})"
        )

    channel_col_idx = find_column_index(
        ws=ws,
        column_name="Channel",
        header_row=header_row,
    )

    url_col_idx = find_column_index(
        ws=ws,
        column_name="URL",
        header_row=header_row,
    )

    if channel_col_idx is None or url_col_idx is None:
        raise ValueError(
            f"{ws.title}: Channel/URL column not found"
        )

    # 본 처리 루프 전에 YT/IG permalink를 채널별로 먼저 훑어서
    # ThreadPoolExecutor로 병렬 prefetch한다. YT/IG는 session/브라우저
    # 상태를 공유하지 않는 순수 yt-dlp 호출이라 병렬화가 안전하지만,
    # X/FB/TT는 Playwright persistent context를 공유해서 이 방식을
    # 적용할 수 없다. 채널별 평균 호출 시간 차이(YT >> IG)에 맞춰
    # max_workers도 채널마다 다르게 준다 (comment_extractor_v2.py 참고).
    loop_start_row = max(
        header_row + 1,
        start_row if start_row is not None else header_row + 1,
    )

    permalinks_to_prefetch: dict[str, list[str]] = {
        "YT": [],
        "IG": [],
    }

    for row_idx in range(loop_start_row, ws.max_row + 1):
        raw_channel = ws.cell(row=row_idx, column=channel_col_idx).value

        if raw_channel is None:
            continue

        normalized_channel = str(raw_channel).strip().upper()

        if normalized_channel not in permalinks_to_prefetch:
            continue

        prefetch_url = extract_post_url(
            ws.cell(row=row_idx, column=url_col_idx).value
        )

        if prefetch_url is None:
            continue

        permalinks_to_prefetch[normalized_channel].append(
            prefetch_url
        )

    for channel, urls in permalinks_to_prefetch.items():
        if not urls:
            continue

        print(
            f"[INFO] {channel} 댓글 URL "
            f"{len(urls)}건 병렬 prefetch 시작..."
        )

        comment_session.prefetch_comment_urls(
            channel,
            urls,
            max_workers=(
                YOUTUBE_PREFETCH_MAX_WORKERS
                if channel == "YT"
                else INSTAGRAM_PREFETCH_MAX_WORKERS
            ),
        )

        print(f"[INFO] {channel} 댓글 URL prefetch 완료")

    processed_count = start_processed_count
    reaction_found_count = start_reaction_found_count

    for row_idx in range(loop_start_row, ws.max_row + 1):
        if checkpoint_saver is not None:
            checkpoint_saver(
                next_row=row_idx,
                processed_count=processed_count,
                reaction_found_count=reaction_found_count,
            )

        raw_channel = ws.cell(row=row_idx, column=channel_col_idx).value
        raw_url = ws.cell(row=row_idx, column=url_col_idx).value

        if raw_channel is None or raw_url is None:
            continue

        post_url = extract_post_url(raw_url)

        if post_url is None:
            continue

        url_cell_value = build_url_cell_value(
            channel=raw_channel,
            post_url=post_url,
            session=comment_session,
            raise_on_error=False,
        )

        if url_cell_value is None:
            continue

        ws.cell(
            row=row_idx,
            column=url_col_idx,
            value=url_cell_value,
        )

        processed_count += 1

        if "\nN/A" not in url_cell_value:
            reaction_found_count += 1

    return processed_count, reaction_found_count


# =========================================================
# 3. 실행 인자
# =========================================================

def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "raw_to_processed.py가 만든 formatted Excel에서 "
            "채널별 소비자 반응 URL을 추출해 URL 컬럼을 갱신합니다."
        )
    )

    parser.add_argument(
        "--output-dir",
        type=Path,
        help=(
            "기존 실행 output 폴더. "
            "예: output/260805_2차. "
            "단독 실행 시 반드시 지정해야 합니다."
        ),
    )

    parser.add_argument(
        "--overwrite",
        action="store_true",
        help=(
            "이 인자는 raw_to_processed.py/media_extractor.py와의 "
            "파이프라인 호출 규격을 맞추기 위해 존재하며, "
            "이 모듈은 항상 같은 formatted Excel을 in-place로 갱신한다."
        ),
    )

    return parser.parse_args()


# =========================================================
# 4. Main Function
# =========================================================

def main() -> None:
    run_started_at = time.monotonic()

    args = parse_arguments()

    input_date = input().strip()

    try:
        datetime.strptime(input_date, "%y%m%d")
    except ValueError as exc:
        raise ValueError(
            "날짜는 YYMMDD 형식으로 입력해야 합니다. "
            "예시: 260724"
        ) from exc

    output_dir = resolve_output_directory(
        input_date=input_date,
        cli_output_dir=args.output_dir,
    )

    formatted_excel_path = build_formatted_excel_path(
        input_date=input_date,
        output_dir=output_dir,
    )

    print(f"Input/Output file: {formatted_excel_path}")

    wait_for_network_ready()

    temporary_path = prepare_temporary_path(formatted_excel_path)

    checkpoint_xlsx_path, checkpoint_json_path = checkpoint_paths(
        formatted_excel_path
    )

    comment_session = CommentExtractorSession()

    excel_fingerprint = source_fingerprint(formatted_excel_path)

    checkpoint_workbook, checkpoint_progress = load_checkpoint(
        checkpoint_xlsx_path,
        checkpoint_json_path,
        excel_fingerprint,
    )

    try:
        if checkpoint_workbook is not None:
            print(
                "[CHECKPOINT] 이전 실행의 중간 저장을 "
                "발견해 이어서 진행합니다: "
                f"{checkpoint_progress}"
            )

            workbook = checkpoint_workbook
            start_row = checkpoint_progress["next_row"]
            start_processed_count = checkpoint_progress[
                "processed_count"
            ]
            start_reaction_found_count = checkpoint_progress[
                "reaction_found_count"
            ]
        else:
            workbook = load_workbook(formatted_excel_path)
            start_row = None
            start_processed_count = 0
            start_reaction_found_count = 0

        if NEW_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"Sheet not found: {NEW_SHEET_NAME} "
                f"(2a 단계 결과가 아닌 것으로 보입니다)"
            )

        target_ws = workbook[NEW_SHEET_NAME]

        checkpoint_saver = make_checkpoint_saver(
            workbook=workbook,
            checkpoint_xlsx_path=checkpoint_xlsx_path,
            checkpoint_json_path=checkpoint_json_path,
            fingerprint=excel_fingerprint,
        )

        processed_count, reaction_found_count = process_reaction_sheet(
            ws=target_ws,
            comment_session=comment_session,
            start_row=start_row,
            start_processed_count=start_processed_count,
            start_reaction_found_count=start_reaction_found_count,
            checkpoint_saver=checkpoint_saver,
        )

        workbook.save(temporary_path)
        os.replace(temporary_path, formatted_excel_path)

        # 정상 완료됐으므로 체크포인트는 더 이상 필요 없다.
        cleanup_checkpoint(checkpoint_xlsx_path, checkpoint_json_path)

    except Exception:
        cleanup_temporary_path(temporary_path)
        raise

    finally:
        timing_csv_path = output_dir / "comment_extraction_timings.csv"

        try:
            comment_session.export_call_records_csv(timing_csv_path)
        except Exception as exc:
            print(
                "[WARNING] 댓글 추출 소요시간 CSV 저장 실패: "
                f"{type(exc).__name__}: {exc}"
            )

        comment_session.close()

    total_elapsed_seconds = time.monotonic() - run_started_at

    print("Done")
    print(f"Processed rows: {processed_count:,}")
    print(f"Consumer reaction URL found: {reaction_found_count:,}")
    print(f"Output file: {formatted_excel_path}")
    print(
        "[TIMING] consumer_reaction_url.py total elapsed: "
        f"{total_elapsed_seconds:.1f}s "
        f"({total_elapsed_seconds / 60:.1f}min)"
    )
    print(f"[TIMING] comment extraction timing CSV: {timing_csv_path}")


if __name__ == "__main__":
    main()
